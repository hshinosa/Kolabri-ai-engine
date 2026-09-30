"""
Text extraction from PDF, DOCX, PPTX, and plain text files.
"""

from __future__ import annotations

import asyncio
import io
import zipfile
import gc
from typing import List, Dict, Any, Optional

from PIL import Image

from app.core.logging import get_logger
from app.services.document_processing.chunking import create_chunks as _create_chunks, ChunkSpec
from app.services.document_processing.models import ProcessedDocument, ProcessedChunk

logger = get_logger(__name__)


def _extract_images_from_docx(doc_content: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(doc_content)) as doc_zip:
            for name in doc_zip.namelist():
                if name.startswith("word/media/"):
                    try:
                        image_data = doc_zip.read(name)
                        img = Image.open(io.BytesIO(image_data))
                        del image_data
                        yield img
                    except Exception:
                        continue
    except Exception as e:
        logger.warning("docx_extract_media_failed", error=str(e))


def _chunks_to_processed(chunks: List[ChunkSpec]) -> List[ProcessedChunk]:
    return [
        ProcessedChunk(text=c.text, metadata=c.metadata, chunk_id=c.chunk_id)
        for c in chunks
    ]


# --- Sync cores: called via asyncio.to_thread so blocking fitz/PIL/zip work
# --- never stalls the event loop during ingest.


def _open_pdf(fitz, content: Optional[bytes], file_path: Optional[str]):
    """Open a PDF from bytes or disk (blocking; run off the event loop)."""
    if file_path is not None:
        return fitz.open(filename=file_path)
    if content is not None:
        return fitz.open(stream=content, filetype="pdf")
    raise ValueError("Either content or file_path must be provided for PDF processing")


def _decode_pdf_image(pdf_doc, xref, min_width: int, min_height: int):
    """Extract and decode one embedded image, None if below the size floor.

    Blocking (archive read + PIL decode); run off the event loop.
    """
    base_image = pdf_doc.extract_image(xref)
    pil_img = Image.open(io.BytesIO(base_image["image"]))
    if pil_img.width < min_width or pil_img.height < min_height:
        pil_img.close()
        return None
    return pil_img


def _parse_docx_text(DocxDocument, content: bytes) -> str:
    """Parse a DOCX into joined paragraph/table text (blocking XML walk)."""
    doc = DocxDocument(io.BytesIO(content))

    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells if cell.text.strip()
            )
            if row_text:
                paragraphs.append(row_text)

    return "\n\n".join(paragraphs)


def _next_image(iterator) -> Any:
    """Pull the next decoded image off a blocking extractor iterator."""
    return next(iterator, None)


def _open_pptx(Presentation, content: bytes):
    """Parse a PPTX from bytes (blocking zip/XML parse; run off the event loop)."""
    return Presentation(io.BytesIO(content))


def _open_pptx_image(shape):
    """Decode a picture shape's blob into a PIL image (blocking; run off the event loop)."""
    return Image.open(io.BytesIO(shape.image.blob))


async def process_pdf(
    content: Optional[bytes],
    filename: str,
    document_id: str,
    metadata: Optional[Dict[str, Any]] = None,
    file_path: Optional[str] = None,
    *,
    ocr_available: bool,
    vision_available: bool,
    min_text_length_for_ocr: int,
    max_images_per_page: int,
    chunk_size: int,
    chunk_overlap: int,
    min_image_width: int = 100,
    min_image_height: int = 100,
    ocr_fn=None,
    caption_fn=None,
    _fitz=None,
):
    fitz = _fitz
    if fitz is None:
        import fitz as _f
        fitz = _f

    chunks: List[ProcessedChunk] = []
    all_text = []
    image_count = 0

    pdf_doc = await asyncio.to_thread(_open_pdf, fitz, content, file_path)
    page_count = len(pdf_doc)
    meta = metadata or {}
    force_ocr = bool(meta.get("perform_ocr"))
    extract_images = meta.get("extract_images", True)
    if extract_images is False:
        extract_images = False
    else:
        extract_images = True

    for page_num, page in enumerate(pdf_doc, start=1):
        page_text_parts = []

        text = await asyncio.to_thread(page.get_text, "text")
        text_length = len(text.strip()) if text else 0

        if text.strip():
            page_text_parts.append(text.strip())

        should_ocr = ocr_available and (force_ocr or text_length < min_text_length_for_ocr)

        if should_ocr and ocr_fn:
            ocr_text = await ocr_fn(page)
            if ocr_text:
                page_text_parts.append(f"[OCR]: {ocr_text.strip()}")
                image_count += 1

        if extract_images and vision_available and caption_fn:
            image_list = page.get_images(full=True)
            for img_index, img in enumerate(image_list[:max_images_per_page]):
                pil_img = None
                try:
                    xref = img[0]
                    pil_img = await asyncio.to_thread(
                        _decode_pdf_image,
                        pdf_doc,
                        xref,
                        min_image_width,
                        min_image_height,
                    )
                    if pil_img is None:
                        continue

                    logger.info(
                        "analyzing_pdf_image", page=page_num, img_index=img_index
                    )
                    caption = await caption_fn(pil_img)

                    if caption:
                        formatted_caption = (
                            f"\n\n=== [GAMBAR VISUAL DI HALAMAN {page_num}] ===\n"
                            f"Deskripsi: {caption}\n"
                            f"=============================================\\n\n"
                        )
                        page_text_parts.append(formatted_caption)
                        image_count += 1
                except Exception as e:
                    logger.warning(
                        "pdf_image_extraction_failed", page=page_num, error=str(e)
                    )
                finally:
                    if pil_img is not None:
                        pil_img.close()
                        del pil_img
                    gc.collect()

        page_text = "\n\n".join(page_text_parts)
        if page_text:
            all_text.append(page_text)
            specs = await asyncio.to_thread(
                _create_chunks,
                text=page_text,
                document_id=document_id,
                filename=filename,
                page_number=page_num,
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
                metadata=metadata,
            )
            chunks.extend(_chunks_to_processed(specs))

    pdf_doc.close()
    gc.collect()

    return ProcessedDocument(
        filename=filename,
        file_type="pdf",
        chunks=chunks,
        page_count=page_count,
        image_count=image_count,
        total_characters=sum(len(c.text) for c in chunks),
        processing_time_ms=0,
        success=True,
    )


async def process_docx(
    content: bytes,
    filename: str,
    document_id: str,
    metadata: Optional[Dict[str, Any]] = None,
    *,
    chunk_size: int,
    chunk_overlap: int,
    ocr_available: bool = False,
    ocr_fn=None,
    _DocxDocument=None,
    _extract_images_fn=None,
):
    DocxDocument = _DocxDocument
    if DocxDocument is None:
        from docx import Document as DocxDocument

    extract_fn = _extract_images_fn or _extract_images_from_docx

    # Blocking XML parse + paragraph/table walk; run off the event loop
    full_text = await asyncio.to_thread(_parse_docx_text, DocxDocument, content)

    image_count = 0
    if ocr_available and ocr_fn:
        try:
            ocr_texts = []
            processed = 0
            images = iter(extract_fn(content))
            while True:
                img = await asyncio.to_thread(_next_image, images)
                if img is None:
                    break
                if processed >= 5:
                    img.close()
                    break
                try:
                    text = await ocr_fn(img)
                    if text:
                        ocr_texts.append(f"[IMAGE_OCR]: {text}")
                finally:
                    img.close()
                    del img
                    gc.collect()
                processed += 1
            if ocr_texts:
                full_text += "\n\n" + "\n\n".join(ocr_texts)
                image_count = len(ocr_texts)
        except Exception as e:
            logger.warning("docx_image_extraction_failed", error=str(e))

    specs = await asyncio.to_thread(
        _create_chunks,
        text=full_text,
        document_id=document_id,
        filename=filename,
        page_number=1,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        metadata=metadata,
    )
    chunks = _chunks_to_processed(specs)

    return ProcessedDocument(
        filename=filename,
        file_type="docx",
        chunks=chunks,
        page_count=1,
        image_count=image_count,
        total_characters=sum(len(c.text) for c in chunks),
        processing_time_ms=0,
        success=True,
    )


async def process_pptx(
    content: bytes,
    filename: str,
    document_id: str,
    metadata: Optional[Dict[str, Any]] = None,
    *,
    chunk_size: int,
    chunk_overlap: int,
    ocr_available: bool = False,
    ocr_fn=None,
    _Presentation=None,
):
    Presentation = _Presentation
    if Presentation is None:
        from pptx import Presentation

    prs = await asyncio.to_thread(_open_pptx, Presentation, content)

    chunks: List[ProcessedChunk] = []
    total_chars = 0
    image_count = 0

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text_parts.append(shape.text.strip())

            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        slide_text_parts.append(row_text)

            if shape.shape_type == 13 and ocr_available and ocr_fn:
                img = None
                try:
                    img = await asyncio.to_thread(_open_pptx_image, shape)
                    ocr_text = await ocr_fn(img)
                    if ocr_text:
                        slide_text_parts.append(f"[IMAGE_OCR]: {ocr_text}")
                        image_count += 1
                except Exception as e:
                    logger.warning("pptx_image_ocr_failed", error=str(e))
                finally:
                    if img is not None:
                        img.close()
                        del img
                    gc.collect()

        slide_text = "\n".join(slide_text_parts)
        if slide_text:
            total_chars += len(slide_text)
            slide_specs = await asyncio.to_thread(
                _create_chunks,
                text=slide_text,
                document_id=document_id,
                filename=filename,
                page_number=slide_num,
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
                metadata={**(metadata or {}), "slide_number": slide_num},
            )
            chunks.extend(_chunks_to_processed(slide_specs))

    return ProcessedDocument(
        filename=filename,
        file_type="pptx",
        chunks=chunks,
        page_count=len(prs.slides),
        image_count=image_count,
        total_characters=total_chars,
        processing_time_ms=0,
        success=True,
    )


def _process_text_sync(
    content: bytes,
    filename: str,
    document_id: str,
    file_type: str,
    metadata: Optional[Dict[str, Any]],
    chunk_size: int,
    chunk_overlap: int,
) -> ProcessedDocument:
    text = None
    for encoding in ["utf-8", "utf-16", "latin-1", "cp1252"]:
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        raise ValueError("Unable to decode text file with supported encodings")

    specs = _create_chunks(
        text=text,
        document_id=document_id,
        filename=filename,
        page_number=1,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        metadata=metadata,
    )
    chunks = _chunks_to_processed(specs)

    return ProcessedDocument(
        filename=filename,
        file_type=file_type,
        chunks=chunks,
        page_count=1,
        image_count=0,
        total_characters=len(text),
        processing_time_ms=0,
        success=True,
    )


async def process_text(
    content: bytes,
    filename: str,
    document_id: str,
    file_type: str,
    metadata: Optional[Dict[str, Any]] = None,
    *,
    chunk_size: int,
    chunk_overlap: int,
):
    # Decode + chunk are pure CPU work; run off the event loop
    return await asyncio.to_thread(
        _process_text_sync,
        content,
        filename,
        document_id,
        file_type,
        metadata,
        chunk_size,
        chunk_overlap,
    )
